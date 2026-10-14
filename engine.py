"""Calculation engine for the ₹1 Cr / 3-year portfolio dashboard.

Reads the raw inputs from the assignment workbook and reproduces the workbook's
model (factor scores -> investor score -> eligibility -> segment allocation ->
expected return -> projection, backtest and probability of target) in pandas,
so every number can be recomputed instantly when an input changes.
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

DEFAULT_XLSX = Path(__file__).parent / "data" / "Investment_Management_Portfolio.xlsx"

FACTORS = [  # (name, column in universe frame, higher_is_better)
    ("Market Cap", "mcap", True),
    ("Sales Growth YoY", "sales_yoy", True),
    ("EPS Growth YoY", "eps_growth", True),
    ("ROE", "roe", True),
    ("ROCE", "roce", True),
    ("Trailing P/E", "pe", False),
    ("PEG", "peg", False),
    ("P/BV", "pb", False),
    ("Beta", "beta", False),
    ("Annualised SD", "sd", False),
    ("Price vs 10M SMA", "vs_sma", True),
    ("6M Momentum", "mom6", True),
    ("Debt / Equity", "de", False),
]
SEGMENTS = ["Large Cap", "Mid Cap", "Small Cap"]


# ----------------------------------------------------------------------------- loading
@dataclass
class ModelData:
    universe: pd.DataFrame          # one row per stock, raw inputs
    prices: pd.DataFrame            # month-end adjusted closes, columns = tickers + NIFTY_50
    debt: pd.DataFrame              # one row per bond
    personas: pd.DataFrame          # one row per investor, policy presets
    factor_weights: pd.DataFrame    # rows = factors, cols = House + investors
    market: dict = field(default_factory=dict)  # rf, erp, blend, tilt, gcap, corpus, horizon


def _f(x, default=np.nan):
    try:
        return float(x)
    except (TypeError, ValueError):
        return default


def load_model(xlsx: str | Path | None = None) -> ModelData:
    wb = load_workbook(xlsx or DEFAULT_XLSX, data_only=True)
    fs, pdw, ip, sm = wb["Master Fund Sheet"], wb["Price_Data"], wb["Investor Personas"], wb["Scoring Model"]

    rows = []
    for r in range(2, 16):
        rows.append(dict(
            ticker=fs[f"C{r}"].value, company=fs[f"D{r}"].value, sector=fs[f"E{r}"].value,
            price=_f(fs[f"I{r}"].value), mcap=_f(fs[f"P{r}"].value), eps_ttm=_f(fs[f"Q{r}"].value),
            roce=_f(fs[f"S{r}"].value), roe=_f(fs[f"T{r}"].value), bv=_f(fs[f"U{r}"].value),
            sales_yoy=_f(fs[f"W{r}"].value), profit_3y=_f(fs[f"Y{r}"].value), de=_f(fs[f"AA{r}"].value),
            dy=_f(fs[f"AB{r}"].value), eps_fy25=_f(fs[f"AG{r}"].value), eps_fy26=_f(fs[f"AH{r}"].value),
        ))
    universe = pd.DataFrame(rows).set_index("ticker")

    debt = pd.DataFrame([dict(
        code=fs[f"C{r}"].value, name=fs[f"D{r}"].value, kind=fs[f"B{r}"].value, rating=fs[f"G{r}"].value,
        unit_price=_f(fs[f"H{r}"].value), ytm=_f(fs[f"I{r}"].value), pd=_f(fs[f"J{r}"].value, 0), lgd=_f(fs[f"K{r}"].value, 0),
    ) for r in (18, 19, 20)]).set_index("code")
    debt["credit_adj"] = debt["ytm"] - debt["pd"] * debt["lgd"]

    cols = [pdw.cell(row=1, column=c).value for c in range(2, 17)]
    data = []
    for r in range(2, 63):
        data.append([pdw.cell(row=r, column=c).value for c in range(1, 17)])
    prices = pd.DataFrame(data, columns=["month"] + cols).set_index("month")
    prices.index = pd.to_datetime(prices.index)
    prices = prices.astype(float)

    # presets: segment split & debt mix live in Dash_Calc (numeric table); fall back to defaults
    seg_mix = {"Vikram": (0.40, 0.35, 0.25, 0.0, 0.0, 1.0), "Ananya": (0.50, 0.30, 0.20, 0.34, 0.33, 0.33),
               "Rohan": (0.65, 0.25, 0.10, 0.43, 0.34, 0.23), "Meera": (0.80, 0.15, 0.05, 0.50, 0.33, 0.17)}
    if "Dash_Calc" in wb.sheetnames:
        dc = wb["Dash_Calc"]
        for j in range(4):
            name = dc.cell(row=40, column=2 + j).value
            if name:
                seg_mix[name] = tuple(_f(dc.cell(row=41 + k, column=2 + j).value, 0) for k in range(6))
    prs = []
    for r in range(3, 7):
        nm = ip[f"A{r}"].value
        L, M, S, g, a, s = seg_mix[nm]
        prs.append(dict(name=nm, archetype=ip[f"B{r}"].value, risk=ip[f"E{r}"].value, target=_f(ip[f"C{r}"].value),
                        equity=_f(ip[f"F{r}"].value), max_beta=_f(ip[f"I{r}"].value), min_score=_f(ip[f"J{r}"].value),
                        profile=ip[f"K{r}"].value, large=L, mid=M, small=S, debt_gsec=g, debt_aaa=a, debt_a=s))
    personas = pd.DataFrame(prs).set_index("name")

    fw = pd.DataFrame(
        [[_f(sm.cell(row=5 + k, column=4 + j).value) for j in range(5)] for k in range(13)],
        index=[f[0] for f in FACTORS], columns=["House", "Vikram", "Ananya", "Rohan", "Meera"])

    market = dict(corpus=_f(ip["C1"].value, 1e7), horizon=_f(ip["F1"].value, 3), erp=_f(ip["C10"].value, 0.07),
                  blend=_f(ip["C11"].value, 0.5), tilt=_f(ip["C12"].value, 0.03), rf=_f(ip["C13"].value, 0.0735),
                  gcap=_f(ip["C14"].value, 0.22))
    return ModelData(universe, prices, debt, personas, fw, market)


# ----------------------------------------------------------------------------- analytics
def enrich(u: pd.DataFrame, prices: pd.DataFrame) -> pd.DataFrame:
    """Add derived ratios, risk metrics and technicals (same definitions as the workbook)."""
    u = u.copy()
    rets = prices.pct_change().iloc[1:]
    mkt = rets["NIFTY_50"]
    u["pe"] = np.where(u["eps_ttm"] > 0, u["price"] / u["eps_ttm"], np.nan)
    u["pb"] = u["price"] / u["bv"]
    u["eps_growth"] = np.where(u["eps_fy25"] > 0, u["eps_fy26"] / u["eps_fy25"] - 1, np.nan)
    u["peg"] = np.where((u["profit_3y"] > 0) & u["pe"].notna(), u["pe"] / (u["profit_3y"] * 100), np.nan)
    beta, sd, cagr5, cagr3, ret1, sma, mom6 = {}, {}, {}, {}, {}, {}, {}
    for t in u.index:
        r = rets[t]
        beta[t] = np.cov(r, mkt, ddof=1)[0, 1] / np.var(mkt, ddof=1)
        sd[t] = r.std(ddof=1) * math.sqrt(12)
        p = prices[t]
        cagr5[t] = (p.iloc[-1] / p.iloc[0]) ** (1 / 5) - 1
        cagr3[t] = (p.iloc[-1] / p.iloc[-37]) ** (1 / 3) - 1
        ret1[t] = p.iloc[-1] / p.iloc[-13] - 1
        sma[t] = p.iloc[-10:].mean()
        mom6[t] = p.iloc[-1] / p.iloc[-7] - 1
    u["beta"], u["sd"], u["cagr5"], u["cagr3"], u["ret1"] = pd.Series(beta), pd.Series(sd), pd.Series(cagr5), pd.Series(cagr3), pd.Series(ret1)
    u["sma10"] = pd.Series(sma)
    u["vs_sma"] = u["price"] / u["sma10"] - 1
    u["mom6"] = pd.Series(mom6)
    u["segment"] = np.where(u["mcap"] >= 100000, "Large Cap", np.where(u["mcap"] >= 30000, "Mid Cap", "Small Cap"))
    return u


def factor_scores(u: pd.DataFrame) -> pd.DataFrame:
    """Percentile score 0-100: (rank-1)/(n-1)*100 in the 'better' direction (Excel RANK ties)."""
    out = {}
    for name, col, higher in FACTORS:
        x = u[col]
        valid = x.dropna()
        n = len(valid)
        rank = valid.rank(ascending=higher, method="min")
        s = (rank - 1) / (n - 1) * 100 if n > 1 else rank * 0 + 50
        default = 50.0 if col == "de" else 0.0          # bank D/E neutral; negative-growth PEG = 0
        out[name] = s.reindex(u.index).fillna(default)
    return pd.DataFrame(out)


@dataclass
class Policy:
    corpus: float
    target: float
    horizon: float
    equity: float
    large: float
    mid: float
    small: float
    max_beta: float
    min_score: float
    debt_gsec: float
    debt_aaa: float
    debt_a: float
    rf: float
    erp: float
    blend: float
    tilt: float
    gcap: float
    weights: pd.Series  # factor weights


def preset_policy(md: ModelData, investor: str) -> Policy:
    p = md.personas.loc[investor]; m = md.market
    return Policy(corpus=m["corpus"], target=p["target"], horizon=m["horizon"], equity=p["equity"],
                  large=p["large"], mid=p["mid"], small=p["small"], max_beta=p["max_beta"], min_score=p["min_score"],
                  debt_gsec=p["debt_gsec"], debt_aaa=p["debt_aaa"], debt_a=p["debt_a"],
                  rf=m["rf"], erp=m["erp"], blend=m["blend"], tilt=m["tilt"], gcap=m["gcap"],
                  weights=md.factor_weights[investor].copy())


def run(md: ModelData, pol: Policy, universe: pd.DataFrame | None = None) -> dict:
    u = enrich(universe if universe is not None else md.universe, md.prices)
    sc = factor_scores(u)
    w = pol.weights.reindex(sc.columns).fillna(0)
    tw = w.sum()
    u["score"] = (sc * w).sum(axis=1) / tw if tw > 0 else 0.0

    # expected return
    capm = pol.rf + u["beta"] * pol.erp
    payout = np.where(u["pe"].notna(), np.minimum(u["dy"] * u["pe"], 1), 1)
    g = np.minimum(np.maximum(u["roe"] * (1 - payout), 0), pol.gcap)
    fund = u["dy"] + g
    u["capm"], u["fundamental"] = capm, fund
    u["exp_return"] = pol.blend * capm + (1 - pol.blend) * fund + (u["score"] - 50) / 50 * pol.tilt

    # allocation
    u["eligible"] = (u["score"] >= pol.min_score) & (u["beta"] <= pol.max_beta)
    share = {"Large Cap": pol.large, "Mid Cap": pol.mid, "Small Cap": pol.small}
    has = {s: bool(((u["segment"] == s) & u["eligible"]).any()) for s in SEGMENTS}
    denom = sum(share[s] for s in SEGMENTS if has[s])
    eff = {s: (share[s] / denom if (has[s] and denom > 0) else 0.0) for s in SEGMENTS}
    seg_sum = u[u["eligible"]].groupby("segment")["score"].sum()
    u["weight"] = [
        pol.equity * eff[seg] * sc_ / seg_sum[seg] if (el and seg_sum.get(seg, 0) > 0) else 0.0
        for seg, sc_, el in zip(u["segment"], u["score"], u["eligible"])
    ]
    d = md.debt.copy()
    mix = np.array([pol.debt_gsec, pol.debt_aaa, pol.debt_a], dtype=float)
    d["weight"] = (1 - pol.equity) * mix / mix.sum() if mix.sum() > 0 else 0.0

    # units / values
    u["units"] = np.floor(u["weight"] * pol.corpus / u["price"])
    u["invested"] = u["units"] * u["price"]
    u["value_h"] = u["invested"] * (1 + u["exp_return"]) ** pol.horizon
    u["target_price"] = u["price"] * (1 + u["exp_return"]) ** pol.horizon
    d["units"] = np.floor(d["weight"] * pol.corpus / d["unit_price"])
    d["invested"] = d["units"] * d["unit_price"]
    d["value_h"] = d["invested"] * (1 + d["credit_adj"]) ** pol.horizon
    invested = u["invested"].sum() + d["invested"].sum()
    cash = pol.corpus - invested
    gsec_y = md.debt["ytm"].iloc[0]
    projected = u["value_h"].sum() + d["value_h"].sum() + cash * (1 + gsec_y) ** pol.horizon
    exp_cagr = (projected / pol.corpus) ** (1 / pol.horizon) - 1
    req_cagr = (pol.target / pol.corpus) ** (1 / pol.horizon) - 1
    beta_p = (u["invested"] * u["beta"]).sum() / pol.corpus

    # hypothetical 5Y backtest with today's weights
    rets = md.prices.pct_change().iloc[1:]
    dw = d["weight"].sum()
    dy = (d["weight"] * d["credit_adj"]).sum() / dw if dw > 0 else 0.0
    port = rets[u.index].mul(u["weight"], axis=1).sum(axis=1) + dw * dy / 12
    idx = 100 * (1 + port).cumprod(); nidx = 100 * (1 + rets["NIFTY_50"]).cumprod()
    vol = port.std(ddof=1) * math.sqrt(12)
    dd = idx / idx.cummax() - 1
    prob = 0.0
    if vol > 0:
        z = (math.log(pol.target / pol.corpus) - math.log(1 + exp_cagr) * pol.horizon) / (vol * math.sqrt(pol.horizon))
        prob = 1 - 0.5 * (1 + math.erf(z / math.sqrt(2)))
    years = np.arange(0, int(round(pol.horizon)) + 1)
    path = pd.DataFrame({
        "Year": years,
        "Expected": pol.corpus * (1 + exp_cagr) ** years,
        "Target path": pol.corpus * (1 + req_cagr) ** years,
        "Downside (−1σ)": pol.corpus * (1 + exp_cagr) ** years * np.exp(-vol * np.sqrt(years)),
        "Upside (+1σ)": pol.corpus * (1 + exp_cagr) ** years * np.exp(vol * np.sqrt(years)),
    })
    seg_mix = {s: u.loc[u["segment"] == s, "invested"].sum() / pol.corpus for s in SEGMENTS}
    seg_mix["Debt"] = d["invested"].sum() / pol.corpus
    seg_mix["Cash"] = cash / pol.corpus
    kpi = dict(invested=invested, cash=cash, projected=projected, exp_cagr=exp_cagr, req_cagr=req_cagr,
               gap=projected - pol.target, beta=beta_p, vol=vol, sharpe=(exp_cagr - pol.rf) / vol if vol > 0 else 0,
               prob=prob, n_stocks=int((u["weight"] > 0).sum()),
               bt_cagr=(idx.iloc[-1] / 100) ** (12 / len(port)) - 1, bt_mdd=dd.min(),
               nifty_cagr=(nidx.iloc[-1] / 100) ** (12 / len(port)) - 1)
    return dict(stocks=u, scores=sc, debt=d, kpi=kpi, path=path, backtest=pd.DataFrame({"Portfolio": idx, "NIFTY 50": nidx}),
                drawdown=dd, seg_mix=seg_mix)


if __name__ == "__main__":
    md = load_model()
    for inv in md.personas.index:
        res = run(md, preset_policy(md, inv))
        k = res["kpi"]
        print(f"{inv:7s} projected ₹{k['projected']:,.0f}  exp {k['exp_cagr']:.2%}  req {k['req_cagr']:.2%}  "
              f"prob {k['prob']:.1%}  beta {k['beta']:.2f}  vol {k['vol']:.1%}  stocks {k['n_stocks']}")
